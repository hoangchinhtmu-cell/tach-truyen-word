import io
import re
import zipfile
from copy import deepcopy
from pathlib import Path

import streamlit as st
from docx import Document

st.set_page_config(page_title="Tách truyện Word", page_icon="📖", layout="centered")

WORD_RE = re.compile(r"\S+")

def word_count(text):
    return len(WORD_RE.findall(text or ""))

def split_sentences(text):
    if not text or not text.strip():
        return []
    # Tách sau dấu câu, giữ nguyên dấu câu.
    parts = re.split(r"(?<=[.!?。！？…])(?=\s+|$)", text.strip())
    return [p for p in parts if p.strip()]

def clone_paragraph(src_p, dst_doc):
    """
    Copy trực tiếp XML của paragraph sang document mới.
    Cách này tránh lỗi IndexError của bản V3 và giữ formatting/runs tốt hơn.
    """
    body = dst_doc._body._element
    new_p = deepcopy(src_p._p)

    # sectPr phải luôn ở cuối body; chèn paragraph ngay trước sectPr.
    sectPr = body.sectPr
    if sectPr is not None:
        body.insert(body.index(sectPr), new_p)
    else:
        body.append(new_p)

def clear_default_paragraph(doc):
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

def make_doc_from_units(units):
    out = Document()
    clear_default_paragraph(out)

    for kind, obj in units:
        if kind == "paragraph":
            clone_paragraph(obj, out)
        else:
            out.add_paragraph(obj)

    return out

def paragraph_units(doc, max_words):
    units = []
    for p in doc.paragraphs:
        text = p.text or ""
        if not text.strip():
            units.append(("paragraph", p, 0))
            continue

        wc = word_count(text)

        if wc <= max_words:
            units.append(("paragraph", p, wc))
        else:
            sentences = split_sentences(text)
            if len(sentences) <= 1:
                units.append(("paragraph", p, wc))
            else:
                for s in sentences:
                    units.append(("text", s, word_count(s)))
    return units

def build_chunks(doc, target, tolerance):
    min_words = max(1, target - tolerance)
    max_words = target + tolerance

    units = paragraph_units(doc, max_words)

    chunks = []
    current = []
    current_words = 0

    for kind, obj, wc in units:
        if wc == 0:
            if current:
                current.append((kind, obj))
            continue

        # Nếu chunk hiện tại đã đạt mức tối thiểu và thêm đơn vị này
        # sẽ vượt max, chốt ở điểm hợp lý trước đó.
        if current and current_words >= min_words and current_words + wc > max_words:
            chunks.append(current)
            current = []
            current_words = 0

        current.append((kind, obj))
        current_words += wc

    if current:
        chunks.append(current)

    return chunks

def chunk_count(chunk):
    total = 0
    for kind, obj in chunk:
        total += word_count(obj.text if kind == "paragraph" else obj)
    return total

def export_chunk(chunk):
    return make_doc_from_units(chunk)

def prepare_zip(chunks):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for idx, chunk in enumerate(chunks, 1):
            doc = export_chunk(chunk)
            dbuf = io.BytesIO()
            doc.save(dbuf)
            z.writestr(f"Chuong_{idx:03d}.docx", dbuf.getvalue())
    return buf.getvalue()

st.title("📖 Tách truyện Word")
st.caption("Chia một file Word thành các chương theo số từ, ưu tiên kết thúc ở cuối đoạn/câu và giữ định dạng.")

uploaded = st.file_uploader("Chọn file truyện .docx", type=["docx"])

if uploaded is None:
    st.info("👆 Hãy tải file Word lên để bắt đầu.")
    st.stop()

try:
    raw = uploaded.getvalue()
    doc = Document(io.BytesIO(raw))
except Exception as e:
    st.error(f"Không thể đọc file Word: {e}")
    st.stop()

total_words = sum(word_count(p.text) for p in doc.paragraphs)
st.success(f"Đã đọc **{uploaded.name}** · khoảng **{total_words:,} từ** · **{len(doc.paragraphs):,} đoạn văn**.")

st.subheader("⚙️ Cài đặt chia chương")

c1, c2 = st.columns(2)
with c1:
    target = st.number_input("Số từ mục tiêu / chương", min_value=100, max_value=50000, value=5000, step=100)
with c2:
    tolerance = st.number_input("Dung sai ± từ", min_value=0, max_value=10000, value=500, step=100)

min_words = max(1, target - tolerance)
max_words = target + tolerance
st.info(f"🎯 Mục tiêu: **{target:,} từ/chương** · Khoảng cho phép: **{min_words:,}–{max_words:,} từ**")

if st.button("🔍 Tạo bản xem trước", use_container_width=True):
    with st.spinner("Đang phân tích và chia thử..."):
        chunks = build_chunks(doc, target, tolerance)
    st.session_state["chunks"] = chunks
    st.session_state["source_name"] = uploaded.name
    st.session_state.pop("zip_bytes", None)

if "chunks" in st.session_state:
    chunks = st.session_state["chunks"]
    counts = [chunk_count(c) for c in chunks]

    st.success(f"Dự kiến **{len(chunks):,} chương** · Trung bình **{sum(counts)/len(counts):,.0f} từ/chương**.")

    rows = []
    for i, wc in enumerate(counts, 1):
        rows.append({
            "Chương": f"{i:03d}",
            "Số từ": f"{wc:,}",
            "Trạng thái": "✓" if min_words <= wc <= max_words else "⚠"
        })
    st.dataframe(rows, use_container_width=True, hide_index=True)

    if st.button("✂️ Tách và tạo ZIP", type="primary", use_container_width=True):
        with st.spinner("Đang tạo file Word và ZIP..."):
            zip_bytes = prepare_zip(chunks)

        st.session_state["zip_bytes"] = zip_bytes
        original = Path(st.session_state.get("source_name", "truyen.docx")).stem
        st.session_state["zip_name"] = f"{original}_da_tach.zip"

        st.success(f"✅ Đã tạo xong **{len(chunks):,} file Word** · {len(zip_bytes)/(1024*1024):.1f} MB")

    if "zip_bytes" in st.session_state:
        st.download_button(
            "⬇️ TẢI FILE ZIP",
            data=st.session_state["zip_bytes"],
            file_name=st.session_state["zip_name"],
            mime="application/zip",
            use_container_width=True,
        )

st.divider()
st.caption("V4: sửa lỗi IndexError khi sao chép paragraph và giữ sectPr đúng vị trí trong DOCX.")
